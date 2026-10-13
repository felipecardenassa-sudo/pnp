"""Shared styles and small helpers for the spend tracker workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY, INK, INK2, MUTED = "1F3864", "0B0B0B", "52514E", "898781"
WHITE, SUBTLE, INPUT, TOTAL, BANNER, LINE = "FFFFFF", "F5F7FA", "FFF2CC", "E8EEF7", "FFF4D6", "D9D9D9"
CH_COLORS = {"Meta": "2A78D6", "Google Ads": "EB6834", "ChatGPT Ads": "1BAF7A", "LinkedIn Ads": "EDA100"}
RED, AMBER, GREEN = "D03B3B", "9A6700", "006300"

AMT = '#,##0;-#,##0;"–"'
AMT2 = '#,##0.00;-#,##0.00;"–"'
PCT = '0%;-0%;"–"'
DATE_F = 'd mmm yyyy'
DAY_F = 'ddd d mmm'
MON_F = 'mmm yyyy'

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNDER = Border(bottom=Side(style="medium", color=NAVY))


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def cell(ws, ref, value=None, fmt=None, bold=False, size=10, color=INK, fill=None, border=True,
         align=None, valign="center", wrap=False, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(size, bold, color, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = BOX
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    return c


def title(ws, ref, text, sub=None, sub_ref=None):
    cell(ws, ref, text, size=16, bold=True, color=NAVY, border=False)
    if sub is not None:
        cell(ws, sub_ref, sub, size=9, color=INK2, border=False, italic=True)


def section(ws, ref, text, span=None):
    c = cell(ws, ref, text, size=11, bold=True, color=NAVY, border=False)
    if span:
        for col in span:
            ws[f"{col}{c.row}"].border = UNDER
    return c


def header(ws, row, start_col, labels, fill=NAVY, color=WHITE, height=30):
    for i, lab in enumerate(labels):
        ref = f"{get_column_letter(start_col + i)}{row}"
        cell(ws, ref, lab, bold=True, color=color, fill=fill, align="center", wrap=True, size=9)
    ws.row_dimensions[row].height = height


def note(ws, ref, text, color=INK2, size=9):
    return cell(ws, ref, text, size=size, color=color, border=False, italic=True)


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v
