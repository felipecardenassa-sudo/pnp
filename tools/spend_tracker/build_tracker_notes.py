"""Presenter notes for the paid media spend tracker: a 2-page, on-brand PDF Felipe uses to walk Raúl and Tereza through it.
Every number is read from the recalculated tracker workbook (the same file build_spend_pdf.py reads), so the notes match
the sheet. Rebuild after a tracker update to refresh the numbers; edit DECISIONS when Tereza answers them.
Usage (from the repo root):
  python3 -I tools/spend_tracker/build_tracker_notes.py PNP_Paid_Media_Spend_Tracker_2026.xlsx OUT_DIR [--pdf]
"""
import datetime, html, os, re, subprocess, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "brand"))
import brand as B  # noqa: E402

E = html.escape
STEM = "Spend_Tracker_Notes_2026"
CHROMIUM = "/opt/pw-browsers/chromium"
PAGES = 2
MINUS = "−"
# Trebuchet MS has no arrow glyph, so arrows are drawn
ARROW = ('<svg class="arr" viewBox="0 0 12 8" aria-hidden="true"><path d="M1 4h9M7 1l3 3-3 3" fill="none" '
         f'stroke="{B.GRAY}" stroke-width="1.3" stroke-linecap="round" stroke-linejoin="round"/></svg>')


# ------------------------------------------------------------------ tracker values (cached by LibreOffice or Google Sheets)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
OUT = sys.argv[2]
calc, mon, week, loc, ag, sett, oth, mx = (wb[n] for n in ("Calc", "Monthly", "Weekly", "By location", "Agency check",
                                                             "Settings", "Other channels", "Meta export"))


def C(r):
    return calc[f"C{r}"].value


upto = C(3)
if not isinstance(upto, datetime.datetime):
    sys.exit("No calculated values in the workbook: recalculate it in LibreOffice (or download it from Google Sheets) first.")
period_start, period_end = C(4), C(5)
gone, weeks_left = C(9), C(16)
lw_end, lw_start, p3_start = C(13), C(14), C(15)
budget, maxi, spent, used = C(27), C(28), C(30), C(33)
even, ahead = C(35), C(36)
pace, forecast, allow_b, allow_m = C(39), C(42), C(43), C(44)
ended_3w = C(49)
checks_failed = C(93)
fx = sett["B39"].value
channels = {calc[f"B{r}"].value: dict(spent=calc[f"D{r}"].value or 0, week=calc[f"G{r}"].value or 0,
                                        share=calc[f"I{r}"].value or 0) for r in range(20, 24)}
lastweek = C(48)
plan_week = budget * 7 / ((period_end - period_start).days + 1)

months = []
for r in range(5, 17):
    m, tot, plan, diff = (mon[f"{c}{r}"].value for c in "AFGH")
    months.append(dict(month=m, total=tot or 0, plan=plan or 0, diff=diff if isinstance(diff, (int, float)) else None))
complete = [m for m in months if m["diff"] is not None]
quarters = []
for q in range(4):
    ms = months[3 * q:3 * q + 3]
    if all(m["diff"] is not None for m in ms):
        quarters.append((f"Q{q + 1}", sum(m["diff"] for m in ms)))
under_since = None
for m in reversed(complete):
    if m["diff"] >= 0:
        break
    under_since = m["month"]
over_months = [m["month"] for m in complete if m["diff"] > 0]
last_month = complete[-1]


def block(sheet, heading):
    """Rows (name, year total, share) of a 'By location' block, found by its heading."""
    top = next(r for r in range(1, sheet.max_row + 1) if sheet[f"A{r}"].value == heading)
    hdr = [c.value for c in sheet[top + 1]]
    yc, sc = hdr.index("Year") + 1, hdr.index("Share") + 1
    rows = []
    for r in range(top + 2, top + 60):
        name = sheet.cell(r, 1).value
        if name == "Total":
            break
        if name:
            rows.append((name, sheet.cell(r, yc).value or 0, sheet.cell(r, sc).value or 0))
    return rows


regions = [x for x in block(loc, "BY REGION") if x[1] > 0]
types = sorted(block(loc, "BY CAMPAIGN TYPE"), key=lambda x: -x[1])

meta_weeks = meta_ok = 0
g_ratio = []
for r in range(5, 200):
    if not isinstance(ag[f"A{r}"].value, datetime.datetime):
        continue
    if isinstance(ag[f"C{r}"].value, (int, float)) and isinstance(ag[f"O{r}"].value, (int, float)):
        meta_weeks += 1
        meta_ok += abs(ag[f"O{r}"].value) <= 5
    if isinstance(ag[f"Q{r}"].value, (int, float)):
        g_ratio.append(ag[f"Q{r}"].value)
g_same = sum(abs(x - 1) < 0.01 for x in g_ratio)
h1 = next((ag[f"E{r}"].value, ag[f"H{r}"].value) for r in range(5, 200) if str(ag[f"A{r}"].value or "").startswith("H1 "))
gbp_note = any("1.28" in str(ag[f"A{r}"].value or "") for r in range(5, 200))

# FR_Alibaba straight from the Meta export
hdr_row = next(r for r in range(1, 6) if "Campaign name" in [c.value for c in mx[r]])
cols = {c.value: i for i, c in enumerate(mx[hdr_row])}
fr_total, fr_days = 0.0, []
for row in mx.iter_rows(min_row=hdr_row + 1, values_only=True):
    if row[cols["Campaign name"]] == "FR_Alibaba Campaign" and (row[cols["Amount spent (USD)"]] or 0) > 0:
        fr_total += row[cols["Amount spent (USD)"]]
        fr_days.append(row[cols["Reporting starts"]])

gpt = [(oth[f"D{r}"].value or 0, oth[f"L{r}"].value or 0) for r in range(5, 205) if oth[f"B{r}"].value == "ChatGPT Ads"]
gpt_eur, gpt_usd = sum(x for x, _ in gpt), sum(y for _, y in gpt)
ended_notes = {sett[f"A{r}"].value: sett[f"D{r}"].value for r in range(44, 54) if sett[f"A{r}"].value}
gpt_window = re.search(r"\d+ \w+ – \d+ \w+", ended_notes.get("Test 1", "") or "")
cut_note = next((week[f"R{r}"].value for r in range(5, 60) if "budget cut" in str(week[f"R{r}"].value or "").lower()), "")
cut = re.search(r"Meta \$(\d+)/day to \$(\d+)/day \(from (\d+ \w+)\); Google cut ([\d and]+ \w+)", cut_note)

# the text below relies on these; stop rather than print something the tracker doesn't say
assert abs(sum(c["spent"] for c in channels.values()) - spent) < 0.01
assert abs(spent + pace * weeks_left - forecast) < 0.01
assert isinstance(checks_failed, (int, float))
assert fr_days and gpt and gpt_window and cut and gbp_note and h1


# ------------------------------------------------------------------ formatting
def usd(x):
    return (MINUS if x < -0.5 else "") + f"${abs(x):,.0f}"


def signed(x):
    return "$0" if abs(x) < 0.5 else ("+" if x > 0 else MINUS) + f"${abs(x):,.0f}"


def pct(x, d=0):
    return f"{x * 100:.{d}f}%"


def day(x):
    return f"{x.day} {x:%b}"


def mname(x):
    return f"{x:%B}"


# ------------------------------------------------------------------ content
meta, google, gpt_ch, linkedin = (channels[k] for k in ("Meta", "Google Ads", "ChatGPT Ads", "LinkedIn Ads"))
top2 = types[:2]
TYPE_WORDS = {"Meta – lead ads, prospecting": "Meta prospecting lead ads", "Google – Search, non-brand": "Google Search non-brand"}
over_q = [(q, d) for q, d in quarters if d > 0]
over_q_txt = (" and ".join(q for q, _ in over_q) + " ran over plan (" + " and ".join(signed(d) for _, d in over_q) + ")") if over_q else ""
over_txt = (f"from {mname(over_months[0])} to {mname(over_months[-1])}"
            if over_months and len(over_months) == (over_months[-1].month - over_months[0].month + 1) else "in " + ", ".join(mname(m) for m in over_months))
less_b = 1 - allow_b / pace
g_off = len(g_ratio) - g_same

opening = [
    "<b>What it is:</b> one Google Sheet that shows how much of the 2026 paid media budget we've spent, how fast we're spending "
    "now and where we land on 31 Dec.",
    "<b>Where the numbers come from:</b> Meta's and Google Ads' own daily exports (each campaign, each day), not the agency's reports. "
    "ChatGPT Ads and LinkedIn Ads are typed in when they run.",
    "<b>Upkeep:</b> 3 steps every Monday; everything else updates by itself, with 6 automatic checks.",
    f"<b>Headline:</b> at the current pace we end about {usd(forecast - budget)} over the {usd(budget)} budget, "
    f"but {usd(maxi - forecast)} under the {usd(maxi)} maximum.",
]

tiles = [
    ("Spent so far", usd(spent), f"{pct(used)} of the {usd(budget)} budget"),
    ("Current pace", f"{usd(pace)}<small> a week</small>", f"last 3 full weeks, {day(p3_start)} – {day(lw_end)}"),
    ("Forecast for 31 Dec", usd(forecast), f"{usd(forecast - budget)} over budget · {usd(maxi - forecast)} under the maximum"),
    (f"To land on {usd(budget)}", f"{usd(allow_b)}<small> a week</small>", f"{pct(less_b)} less than now"),
]

tabs = [
    ("Dashboard", "The year on one screen: budget, spent, pace and forecast; budget check; spend by channel, month and quarter; "
                  "year-end outlook; regions and markets.",
     f"Meta {usd(meta['spent'])} ({pct(meta['share'])}), Google Ads {usd(google['spent'])} ({pct(google['share'])}), "
     f"ChatGPT Ads {usd(gpt_ch['spent'])}. Red = over budget or over plan; green = within."),
    ("Weekly", f"Each week (Monday to Sunday): spend by channel against the plan ({usd(plan_week)} a week), Meta registrations and "
               "Google conversions with their cost, notes.",
     f"Last full week ({day(lw_start)} – {day(lw_end)}): {usd(lastweek)} (Meta {usd(meta['week'])}, Google Ads {usd(google['week'])}). "
     f"Agency notes until W/C {day(lw_start)}; mine from W/C {day(lw_start + datetime.timedelta(days=7))}."),
    ("Monthly", "Each month by channel against the plan, the running total, 'as billed' columns for the invoice check, clicks and "
                "conversions with their costs.",
     f"Over plan {over_txt}; under plan every month since {mname(under_since)}. "
     f"{mname(last_month['month'])}: {usd(last_month['total'])}, {usd(-last_month['diff'])} under plan." if under_since else
     f"{mname(last_month['month'])}: {usd(last_month['total'])} against a plan of {usd(last_month['plan'])}."),
    ("By location", "Spend by region, market and campaign type, read from the location code at the start of each campaign name.",
     ", ".join(f"{E(n)} {usd(v)} ({pct(s)})" for n, v, s in regions if s >= 0.01) + ". "
     f"{TYPE_WORDS.get(top2[0][0], top2[0][0])} and {TYPE_WORDS.get(top2[1][0], top2[1][0])} are "
     f"{pct((top2[0][1] + top2[1][1]) / spent)} of all spend."),
    ("Agency check", "The agency's weekly Notion reports next to the platform exports, week by week (reference only).",
     f"Meta matched within $5 in {meta_ok} of {meta_weeks} weeks; Google in only {g_same} of {len(g_ratio)}. "
     f"H1: agency {usd(h1[0])}, platforms {usd(h1[1])}."),
    ("Update, Settings, exports", "Update: the Monday routine and 6 checks. Settings: budget, exchange rate, ended campaigns, "
                                  "location codes. Export tabs: the raw data.",
     ("All 6 checks pass" if checks_failed == 0 else f"{checks_failed} check(s) to fix") + f" (data to {day(upto)}). "
     "Only the yellow cells are ever typed in."),
]

how = [
    f"<b>USD everywhere.</b> Meta bills in USD. Google Ads and ChatGPT Ads bill in EUR, converted at 1 EUR = {fx:.4f} USD "
    "(live rate on 13 Oct). Finance's monthly rates can replace it on Settings.",
    f"<b>Data date</b> = the last day in both exports ({day(upto)}).",
    f"<b>Plan</b> = the {usd(budget)} spread evenly over the year: {usd(plan_week)} a week.",
    f"<b>Current pace</b> = the last 3 full weeks ÷ 3, leaving out campaigns that ended for good (FR_Alibaba: {usd(ended_3w)} in those "
    "weeks; the ChatGPT test). Their spend still counts in 'spent so far'.",
    f"<b>Forecast</b> = spent so far + current pace × {weeks_left:.1f} weeks left.",
    f"<b>Market and campaign type</b> come from each campaign name. A new location code is added once on Settings {ARROW} Locations; "
    "the Update tab flags any spend without a market.",
]

monday = [
    f"Meta Ads Manager {ARROW} campaigns by day, 1 Jan to yesterday {ARROW} .csv {ARROW} 'Meta export' tab: replace the data.",
    f"Google Ads {ARROW} saved spend report, 1 Jan to yesterday {ARROW} .csv {ARROW} 'Google export' tab: replace the data.",
    f"Update tab says 'All good' (6 checks) {ARROW} 1–2 lines on last week in the Weekly notes (launches, pauses, budget changes).",
]
monday_extra = ("<b>Only when needed:</b> ChatGPT or LinkedIn spend (Other channels); an ended campaign, a new location code or "
                "Finance's rates (Settings). <b>First Monday of the month:</b> 'as billed' against the invoices. The weekly paid report "
                "uses the same two exports.")

qa = [
    (f"Why is {pct(used)} spent when {pct(gone)} of the year has gone?",
     over_q_txt + (f"; every month since {mname(under_since)} has been under plan" if under_since else "")
     + f". We are still {usd(ahead)} ahead of an even pace ({usd(even)} by {day(upto)})."),
    ("Why did spend drop in September?",
     f"Budgets were cut: Meta from ${cut.group(1)} to ${cut.group(2)} a day from {cut.group(3)}, Google on {cut.group(4)} "
     "(agency notes on the Weekly tab). FR_Alibaba stopped on " + day(max(fr_days)) + " and is left out of the pace."),
    ("Why not use the agency's reports?",
     f"Meta matched, Google didn't: in {g_off} of {len(g_ratio)} weeks its Google figure was converted to dollars, doubled or lower "
     f"than the platform, and the July report used £1 = $1.28 on a euro account. H1: agency {usd(h1[0])}, platforms {usd(h1[1])}."),
    ("How do we know it's right?",
     "6 checks on the Update tab: columns found, dates read, both exports end on the same day, every campaign has a market, the exports "
     "fit, ended campaigns found" + (" (all pass today)." if checks_failed == 0 else ".") + " Once a month: 'as billed' against the invoices."),
    ("LinkedIn and ChatGPT?",
     (f"LinkedIn: no spend in {period_start.year}. " if linkedin["spent"] == 0 else f"LinkedIn: {usd(linkedin['spent'])} so far. ")
     + f"ChatGPT: one 10-day test, €{gpt_eur:,.2f} ({usd(gpt_usd)}). Both are typed on 'Other channels' when they run."),
]

DECISIONS = [
    ("FR_Alibaba", f"Does it count toward this budget? Meta campaign run by Plug and Play's internal team: {usd(fr_total)} from "
                   f"{day(min(fr_days))} to {day(max(fr_days))}."),
    ("ChatGPT Ads", f"Does it count? The 10-day test, {gpt_window.group(0)}: €{gpt_eur:,.2f} ({usd(gpt_usd)})."),
    ("Year-end target", f"Land on {usd(budget)} (≈ {usd(allow_b)} a week, {pct(less_b)} less than now) or keep the current pace "
                        f"(≈ {usd(forecast)}, inside the {usd(maxi)} maximum)?"),
    ("Exchange rate", f"Keep {fx:.4f} or use Finance's monthly rates?"),
    ("2027", "What is the 2027 budget, and when will it be confirmed?"),
]

outlook = [("…at the current pace", pace, forecast), ("…to land on the budget", allow_b, budget), ("…to use the maximum", allow_m, maxi)]


# ------------------------------------------------------------------ HTML
CSS = B.BASE_CSS + f"""
@page {{ size: A4; margin: 0; }}
* {{ box-sizing: border-box; }}
body {{ font-size: 8.9pt; line-height: 1.36; }}
.page {{ width: 210mm; min-height: 296mm; padding: 9.5mm 11mm 6.5mm; display: flex; flex-direction: column; gap: 3mm; break-after: page; }}
.page:last-child {{ break-after: auto; }}
.kicker {{ color: {B.GRAY}; font-size: 7.6pt; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; }}
h1 {{ color: {B.DARK_BLUE}; font-size: 20pt; font-weight: 400; margin: 1px 0 0; line-height: 1.12; }}
.lede {{ color: {B.GRAY}; font-size: 8.4pt; margin-top: 2px; }}
.card {{ padding: 3.1mm 3.8mm; }}
h2 {{ color: {B.DARK_BLUE}; font-size: 10.8pt; margin: 0 0 2mm; font-weight: 700; }}
h2 .hint {{ color: {B.GRAY}; font-size: 7.6pt; font-weight: 400; margin-left: 6px; }}
.y {{ width: 30px; height: 3px; border-radius: 2px; background: {B.YELLOW}; margin-bottom: 5px; }}
ul.list {{ list-style: none; padding: 0; margin: 0; }}
ul.list li {{ position: relative; padding-left: 13px; margin: 0 0 1.5mm; }}
ul.list li:last-child {{ margin-bottom: 0; }}
ul.list li::before {{ content: ""; position: absolute; left: 1px; top: 5px; width: 6px; height: 6px; border-radius: 50%; background: {B.DARK_BLUE}; }}
b {{ color: {B.DARK_BLUE}; }}
.tiles {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 2.8mm; }}
.tile {{ padding: 3mm 3.4mm; }}
.tile .l {{ color: {B.GRAY}; font-size: 8pt; font-weight: 700; }}
.tile .v {{ color: {B.DARK_BLUE}; font-size: 18pt; font-weight: 700; line-height: 1.15; margin-top: 2px; }}
.tile .v small {{ font-size: 9pt; font-weight: 400; color: {B.GRAY}; }}
.tile .c {{ font-size: 7.8pt; margin-top: 2px; }}
.meter-wrap {{ position: relative; padding: 5mm 0 5.5mm; margin-bottom: 1.8mm; }}
.meter {{ position: relative; height: 12px; border-radius: 6px; background: {B.TINT}; }}
.meter .fill {{ position: absolute; left: 0; top: 0; bottom: 0; border-radius: 6px; background: {B.DARK_BLUE}; }}
.meter .mark {{ position: absolute; top: -4px; bottom: -4px; width: 2.5px; margin-left: -1.25px; border-radius: 1px; background: {B.YELLOW}; }}
.meter-wrap .top {{ position: absolute; top: 0; font-size: 7.8pt; font-weight: 700; color: {B.DARK_BLUE}; transform: translateX(-100%); white-space: nowrap; }}
.meter-wrap .bot {{ position: absolute; bottom: 0; font-size: 7.8pt; color: {B.TEXT}; transform: translateX(-50%); white-space: nowrap; }}
.meter-wrap .ends {{ position: absolute; bottom: 0; font-size: 7.2pt; color: {B.GRAY}; }}
.so {{ margin: 0; }}
table.t {{ width: 100%; border-collapse: collapse; }}
table.t th {{ text-align: left; color: {B.GRAY}; font-weight: 700; font-size: 7.6pt; border-bottom: 1.3px solid {B.DARK_BLUE}; padding: 2px 5px 3px; }}
table.t td {{ border-bottom: 1px solid {B.TINT}; padding: 1.6mm 5px; vertical-align: top; }}
table.t tr:last-child td {{ border-bottom: none; }}
table.t td.tab {{ width: 19%; font-weight: 700; color: {B.DARK_BLUE}; }}
table.t td.tab span {{ display: inline-block; width: 15px; height: 15px; line-height: 15px; border-radius: 50%; background: {B.DARK_BLUE};
  color: #fff; font-size: 7.4pt; text-align: center; margin-right: 5px; }}
table.t td.what {{ width: 45%; }}
table.t td.say {{ color: {B.DARK_BLUE}; }}
table.n td, table.n th {{ text-align: right; white-space: nowrap; }}
table.n td:first-child, table.n th:first-child {{ text-align: left; }}
table.n td {{ padding: 1.4mm 5px; }}
table.n tr.now td {{ background: {B.BG}; font-weight: 700; color: {B.DARK_BLUE}; }}
.cols {{ display: grid; grid-template-columns: 1fr 1fr; gap: 3.2mm; }}
ol.steps {{ margin: 0; padding-left: 17px; }}
ol.steps li {{ margin: 0 0 1.5mm; padding-left: 2px; }}
ol.steps li::marker {{ color: {B.DARK_BLUE}; font-weight: 700; }}
.extra {{ font-size: 8.1pt; margin-top: 2.4mm; padding-top: 2mm; border-top: 1px solid {B.TINT}; }}
dl.qa {{ margin: 0; display: grid; grid-template-columns: 1fr 1fr; column-gap: 6mm; row-gap: 2.6mm; }}
dl.qa div {{ break-inside: avoid; }}
dl.qa dt {{ color: {B.DARK_BLUE}; font-weight: 700; margin-bottom: 0.6mm; }}
dl.qa dd {{ margin: 0; }}
.dec {{ display: grid; grid-template-columns: 1.05fr 1fr; gap: 6mm; }}
.dec ol li b {{ display: inline; }}
.arr {{ width: 10px; height: 7px; vertical-align: 0; margin: 0 1px; }}
.foot {{ margin-top: auto; color: {B.GRAY}; font-size: 7pt; display: flex; justify-content: space-between; }}
"""


def head(kicker, title, lede):
    return f'<div><div class="kicker">{kicker}</div><h1>{title}</h1><div class="lede">{lede}</div></div>'


def page(inner, n):
    return (f'<section class="page">{inner}<div class="foot"><span>Plug and Play · Spend tracker · notes for the walk-through · '
            f'platform data to {day(upto)} {upto.year}</span><span>{n} / {PAGES}</span></div></section>')


def meter():
    u, g = min(used, 1.0) * 100, gone * 100
    return (f'<div class="meter-wrap"><div class="top" style="left:{u:.2f}%">{pct(used)} of the budget used</div>'
            f'<div class="meter"><div class="fill" style="width:{u:.2f}%"></div><div class="mark" style="left:{g:.2f}%"></div></div>'
            f'<div class="ends" style="left:0">$0</div><div class="bot" style="left:{g:.2f}%">{pct(gone)} of the year gone ({day(upto)})</div>'
            f'<div class="ends" style="right:0">{usd(budget)}</div></div>')


def build_html():
    p1 = head("Plug and Play · Paid media · my notes", "Spend tracker: walk-through for Raúl and Tereza",
              f"Platform data to {day(upto)} {upto.year} · all amounts in USD · page 1: what to say and show · page 2: questions, "
              "routine, decisions")
    p1 += ('<div class="card"><div class="y"></div><h2>Open with this</h2><ul class="list">'
           + "".join(f"<li>{x}</li>" for x in opening) + "</ul></div>")
    p1 += ('<div class="tiles">' + "".join(
        f'<div class="card tile"><div class="l">{lab}</div><div class="v">{val}</div><div class="c">{sub}</div></div>'
        for lab, val, sub in tiles) + "</div>")
    p1 += (f'<div class="card"><h2>Budget used vs the calendar</h2>{meter()}'
           f'<p class="so">We are {usd(ahead)} ahead of an even pace ({usd(even)} by {day(upto)}): {over_q_txt}'
           + (f"; every month since {mname(under_since)} has been under plan." if under_since else ".") + "</p></div>")
    p1 += ('<div class="card"><h2>Show the tabs in this order<span class="hint">left: what it shows · right: what to point out</span></h2>'
           '<table class="t"><tr><th>Tab</th><th>What it shows</th><th>Point out</th></tr>'
           + "".join(f'<tr><td class="tab"><span>{i}</span>{E(t)}</td><td class="what">{E(w)}</td><td class="say">{s}</td></tr>'
                     for i, (t, w, s) in enumerate(tabs, 1)) + "</table></div>")

    p2 = head("Plug and Play · Paid media · my notes", "Answers, routine and what to ask",
              "Keep this page open for questions. The last card is how to close the meeting.")
    p2 += ('<div class="card"><h2>If they ask…</h2><dl class="qa">'
           + "".join(f"<div><dt>{q}</dt><dd>{a}</dd></div>" for q, a in qa) + "</dl></div>")
    p2 += ('<div class="cols"><div class="card"><h2>How the numbers are calculated</h2><ul class="list">'
           + "".join(f"<li>{x}</li>" for x in how) + "</ul></div>"
           '<div class="card"><h2>Every Monday<span class="hint">the Update tab, 3 steps</span></h2><ol class="steps">'
           + "".join(f"<li>{x}</li>" for x in monday) + f'</ol><div class="extra">{monday_extra}</div></div></div>')
    rows = "".join(
        f'<tr class="{"now" if i == 0 else ""}"><td>{lab}</td><td>{usd(pw)}</td><td>{usd(ye)}</td><td>{signed(ye - budget)}</td>'
        f'<td>{signed(ye - maxi)}</td></tr>' for i, (lab, pw, ye) in enumerate(outlook))
    p2 += ('<div class="card"><div class="y"></div><h2>Close with: what we need from you</h2><div class="dec">'
           '<ol class="steps">' + "".join(f"<li><b>{E(k)}:</b> {v}</li>" for k, v in DECISIONS) + "</ol>"
           '<div><table class="t n"><tr><th>If we spend from now on…</th><th>Per week</th><th>31 Dec</th><th>vs budget</th>'
           f'<th>vs maximum</th></tr>{rows}</table><p class="so" style="margin-top:2mm;font-size:7.9pt">'
           f'{weeks_left:.1f} weeks left to 31 Dec. The Dashboard\'s "Year-end outlook" shows the same figures.</p></div></div></div>')
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Spend tracker notes</title>"
            f"<style>{CSS}</style></head><body>{page(p1, 1)}{page(p2, 2)}</body></html>")


def main():
    os.makedirs(OUT, exist_ok=True)
    html_path = os.path.join(OUT, STEM + ".html")
    open(html_path, "w", encoding="utf-8").write(build_html())
    print("written", html_path)
    print(f"facts: spent {spent:,.2f} · used {used:.4f} · gone {gone:.4f} · pace {pace:,.2f} · forecast {forecast:,.2f} · "
          f"to budget {allow_b:,.2f} · to max {allow_m:,.2f} · ahead {ahead:,.2f} · FR_Alibaba {fr_total:,.2f} "
          f"({min(fr_days):%d %b}–{max(fr_days):%d %b}) · ChatGPT €{gpt_eur:,.2f} = ${gpt_usd:,.2f} · agency Meta {meta_ok}/{meta_weeks}, "
          f"Google {g_same}/{len(g_ratio)} · H1 {h1[0]:,.2f} vs {h1[1]:,.2f}")
    if "--pdf" in sys.argv:
        font = B.font_in_use()
        pdf = os.path.join(OUT, STEM + ".pdf")
        subprocess.run([CHROMIUM, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf}", "file://" + os.path.abspath(html_path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
        pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
        fonts = subprocess.run(["pdffonts", pdf], capture_output=True, text=True).stdout
        names = {re.sub(r"^[A-Z]{6}\+", "", ln.split()[0]) for ln in fonts.splitlines()[2:] if ln.strip()}
        other = sorted(n for n in names if font.replace(" ", "").lower() not in n.lower())
        print(f"PDF: {pdf} · {pages} pages · font {font}"
              + (f" · WARNING: {pages} pages, expected {PAGES}" if pages != PAGES else "")
              + (f" · WARNING non-brand fonts: {other}" if other else ""))


if __name__ == "__main__":
    main()
