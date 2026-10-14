"""ICP vs conversions – Raúl's exercise: break down the marketing email database and the conversions against the ICP.
  JTP Contact us: "Type of profile form" + Original source + Job title
  Content downloads: "RFI Type of profile" + Original source

Builds a branded PDF (no logo) and a Google-Sheets-ready workbook from one aggregates-only data file.
Usage (from the repo root):
  python3 -I tools/icp_report/build_icp_report.py tools/icp_report/icp_data_2026-10-08.json OUT_DIR [--pdf]
The data file holds aggregated counts only (counts of 1-4 already shown as "<5"); raw contact data is never stored.
icp_data_2026-10-08.json comes from the HubSpot read-only pull of 8 Oct 2026 (the first ICP workbook plus the HubSpot
findings workbook). A new pull needs a new data file with the same keys; the asserts below check that the parts add up.
"""
import html, json, os, re, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "brand"))
import brand as B  # noqa: E402

E = html.escape
MINUS = "−"
STEM = "ICP_vs_Conversions_2026"
CHROMIUM = "/opt/pw-browsers/chromium"


# ------------------------------------------------------------------ numbers
def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def fn(v):
    """Count as shown: '<5' and '–' stay as they are."""
    if isinstance(v, str):
        return v
    return "–" if v is None else f"{v:,.0f}"


def words(v):
    """A count inside a sentence: '<5' reads as 'fewer than 5'."""
    return "fewer than 5" if v == "<5" else fn(v)


def share(a, b):
    a, b = num(a), num(b)
    return a / b if a is not None and b else None


def fp(x, d=None):
    if x is None:
        return "–"
    if isinstance(x, str):
        return x
    v = x * 100
    if d is None:
        d = 0 if v >= 10 else 1
    return f"{v:.{d}f}%"


def rate(row, key):
    v = row.get(key)
    return v if isinstance(v, (int, float)) else None


def row_by(rows, key, value):
    for r in rows:
        if r.get(key) == value:
            return r
    raise KeyError(value)


# ------------------------------------------------------------------ data
D = json.load(open(sys.argv[1], encoding="utf-8"))
OUT = sys.argv[2]
S = D["summary"]
J, C, DB = D["jtp"], D["content"], D["database"]


def g(key, grp):
    return S[key][grp]


def facts():
    f = {}
    for grp, d in (("jtp", J), ("content", C)):
        n = g("Contacts", grp)
        prof = {r["Profile"]: r for r in d["by_profile"]}
        icp = {r["ICP result"]: r for r in d["by_icp"]}
        corp = sum(num(prof[p]["Contacts"]) or 0 for p in prof if p.startswith("Corporate"))
        f[grp] = dict(
            contacts=n, submissions=g("Form submissions", grp), forms=g("Forms", grp),
            corporate=corp, corp10k=g("Corporate, 10,000+ employees (any role)", grp),
            strict=g("Strict ICP – Corporate (role + 10,000+)", grp), gov=g("ICP – Government", grp),
            startup=g("Startup profile (secondary ICP group)", grp),
            founders=num(icp["Secondary ICP – Startup founder"]["Contacts"]),
            qualified=g("Qualified / converted by sales", grp), disqualified=g("Disqualified by sales", grp),
            deals=g("Contacts with a deal", grp), customers=g("Became customer", grp),
            no_title=g("Job title empty", grp), size_unknown=g("Company size unknown", grp),
            marketing=sum(num(r["Marketing contacts"]) or 0 for r in d["by_profile"]),
            prof=prof, icp=icp,
        )
        x = f[grp]
        x["deal_rate"] = share(x["deals"], n)
        x["qual_rate"] = share(x["qualified"], n)
        c10 = prof["Corporate – 10,000+ employees"]
        x["c10_deal"] = rate(c10, "With deal %")
        x["c10_qual"] = rate(c10, "Qualified %")
        x["c10_customers"] = c10["Became customer"]
        # the 10,000+ corporates split by role: ICP (right role) / other role / job title missing
        x["c10_other_role"] = num(icp["Not ICP – Corporate (role or size)"]["Corporate, 10,000+ employees"])
        x["c10_no_title"] = num(icp["Partial – Corporate (data missing)"]["Corporate, 10,000+ employees"])
    f["both"] = g("Contacts in both groups", "jtp")
    f["unique"] = f["jtp"]["contacts"] + f["content"]["contacts"] - f["both"]
    src = {r["Original source"]: r for r in C["by_source"]}
    f["content"]["paid_social"] = src["Paid social"]
    f["content"]["offline"] = src["Offline sources"]
    f["jtp"]["src"] = {r["Original source"]: r for r in J["by_source"]}
    f["content"]["src"] = src
    noise = sum(num(f["content"]["prof"][p]["Contacts"]) for p in ("Other", "Job seeker", "Student"))
    f["content"]["noise"] = noise
    # database
    f["db"] = DB
    srcs = DB["options"]["Original Source"]
    f["db_offline"] = next(o["contacts"] for o in srcs if o["option"] == "Offline Sources")
    f["db_sources_total"] = sum(o["contacts"] for o in srcs)
    return f


F = facts()
JF, CF = F["jtp"], F["content"]

# sanity checks: the parts must add up to the totals shown
assert sum(num(r["Contacts"]) or 0 for r in J["by_source"]) == JF["contacts"]
assert sum(num(r["Contacts"]) or 0 for r in C["by_source"]) == CF["contacts"]
assert JF["strict"] + JF["c10_other_role"] + JF["c10_no_title"] == JF["corp10k"]
assert CF["strict"] + CF["c10_other_role"] + CF["c10_no_title"] == CF["corp10k"]
assert sum(num(r["Contacts"]) or 0 for r in C["by_profile"]) == CF["contacts"]

# ------------------------------------------------------------------ labels in plain words
PROFILE_ORDER = ["Corporate – 10,000+ employees", "Government", "Corporate – smaller / unknown size", "Startup",
                 "Investor / VC", "University", "Job seeker", "Student", "Other", "Unknown"]
PROFILE_LABEL = {"Corporate – 10,000+ employees": "Corporate, 10,000+ employees",
                 "Corporate – smaller / unknown size": "Corporate, smaller or size unknown",
                 "Government": "Government", "Startup": "Startup", "Investor / VC": "Investor or VC",
                 "University": "University", "Job seeker": "Job seeker", "Student": "Student", "Other": "Other",
                 "Unknown": "No answer"}
ROLE_LABEL = {"Empty": "No job title", "C-level": "C-level (CEO, CTO, MD…)", "Other title": "Other title",
              "Founder": "Founder", "Innovation / R&D / Digital": "Innovation, R&D or digital (ICP role)",
              "Government / policy": "Government or policy"}
ICP_KIND = {"Corporate – 10,000+ employees": "icp", "Government": "icp", "Startup": "sec"}
ICP_ORDER = ["ICP – Corporate", "ICP – Government", "Secondary ICP – Startup founder", "Partial – Corporate (data missing)",
             "Partial – Startup (no job title)", "Not ICP – Corporate (role or size)", "Not ICP – Startup (other role)",
             "Not ICP – Other profile"]
ICP_LABEL = {"ICP – Corporate": "ICP corporate: ICP role, 10,000+ employees, $100M+",
             "ICP – Government": "ICP government",
             "Secondary ICP – Startup founder": "Secondary ICP: startup founder or C-level",
             "Partial – Corporate (data missing)": "Maybe ICP: corporate, job title or size missing",
             "Partial – Startup (no job title)": "Maybe secondary ICP: startup, no job title",
             "Not ICP – Corporate (role or size)": "Not ICP: corporate, other role, under 10,000 or under $100M",
             "Not ICP – Startup (other role)": "Not ICP: startup, other role",
             "Not ICP – Other profile": "Not ICP: investor, university, student, job seeker, other, no answer"}


# ------------------------------------------------------------------ HTML pieces
def bar(value, maximum, kind="", width=100):
    v = num(value)
    w = 0 if not v or not maximum else max(1.5, 100 * v / maximum)
    return f'<span class="bar {kind}" style="width:{width}px"><i style="width:{w:.1f}%"></i></span>'


def flag(x, avg, worse_if_higher, contacts=None):
    """Wine when clearly worse than the group average, green when clearly better (used on rates).
    Rows under 100 contacts are never coloured: their rates move a lot with a few people."""
    if x is None or avg is None or avg == 0 or (num(contacts) or 0) < 100:
        return ""
    ratio = x / avg
    if worse_if_higher:
        return "bad" if ratio >= 1.8 else ("good" if ratio <= 0.5 else "")
    return "good" if ratio >= 1.8 else ("bad" if ratio <= 0.4 else "")


def profile_table(d, f):
    rows = {r["Profile"]: r for r in d["by_profile"]}
    order = [p for p in PROFILE_ORDER if p in rows]
    mx = max(num(rows[p]["Contacts"]) or 0 for p in order)
    n = f["contacts"]
    avg_q, avg_dq, avg_deal = share(f["qualified"], n), share(f["disqualified"], n), share(f["deals"], n)
    out = ['<table class="t"><tr><th>Type of profile</th><th class="l">Contacts</th><th></th><th>Share</th>'
           '<th>Qualified by sales</th><th>Disqualified by sales</th><th>With a deal</th></tr>']
    for p in order:
        r = rows[p]
        kind = ICP_KIND.get(p, "oth")
        q, dl = rate(r, "Qualified %"), rate(r, "With deal %")
        tag = '<span class="tag icp">ICP</span>' if kind == "icp" else ('<span class="tag sec">2nd ICP</span>' if kind == "sec" else "")
        out.append(f'<tr><td>{E(PROFILE_LABEL[p])} {tag}</td><td class="l">{bar(r["Contacts"], mx, kind)}</td>'
                   f'<td>{fn(r["Contacts"])}</td><td>{fp(share(r["Contacts"], n))}</td>'
                   f'<td class="{flag(q, avg_q, False, r["Contacts"])}">{fp(r.get("Qualified %"), 1)}</td>'
                   f'<td>{fp(r.get("Disqualified %"), 1)}</td>'
                   f'<td class="{flag(dl, avg_deal, False, r["Contacts"])}">{fp(r.get("With deal %"), 1)}</td></tr>')
    out.append(f'<tr class="tot"><td>All contacts</td><td></td><td>{fn(n)}</td><td>100%</td><td>{fp(avg_q, 1)}</td>'
               f'<td>{fp(avg_dq, 1)}</td><td>{fp(avg_deal, 1)}</td></tr></table>')
    return "".join(out)


def source_table(d, f):
    rows = d["by_source"]
    mx = max(num(r["Contacts"]) or 0 for r in rows)
    n = f["contacts"]
    avg_q, avg_dq, avg_deal = share(f["qualified"], n), share(f["disqualified"], n), share(f["deals"], n)
    avg_c10 = share(f["corp10k"], n)
    out = ['<table class="t"><tr><th>Original source</th><th class="l">Contacts</th><th></th><th>Share</th>'
           '<th>Corporates 10,000+</th><th>Qualified by sales</th><th>Disqualified by sales</th><th>With a deal</th></tr>']
    for r in rows:
        c10 = share(r["Corporate, 10,000+ employees"], r["Contacts"])
        q, dl = rate(r, "Qualified %"), rate(r, "With deal %")
        out.append(f'<tr><td>{E(r["Original source"])}</td><td class="l">{bar(r["Contacts"], mx, "src")}</td>'
                   f'<td>{fn(r["Contacts"])}</td><td>{fp(share(r["Contacts"], n))}</td>'
                   f'<td class="{flag(c10, avg_c10, False, r["Contacts"])}">{fp(c10, 1) if num(r["Corporate, 10,000+ employees"]) is not None else "–"}</td>'
                   f'<td class="{flag(q, avg_q, False, r["Contacts"])}">{fp(r.get("Qualified %"), 1)}</td>'
                   f'<td>{fp(r.get("Disqualified %"), 1)}</td>'
                   f'<td class="{flag(dl, avg_deal, False, r["Contacts"])}">{fp(r.get("With deal %"), 1)}</td></tr>')
    out.append(f'<tr class="tot"><td>All contacts</td><td></td><td>{fn(n)}</td><td>100%</td><td>{fp(avg_c10, 1)}</td>'
               f'<td>{fp(avg_q, 1)}</td><td>{fp(avg_dq, 1)}</td><td>{fp(avg_deal, 1)}</td></tr></table>')
    return "".join(out)


LEGEND = ('<div class="legend">Green = clearly better than the average for all contacts, wine = clearly worse (groups of 100+ contacts). '
          'Disqualified is highest where sales actually reviews the leads, so it shows sales attention as much as lead quality. '
          '"&lt;5" = 1 to 4 contacts; "–" = no rate under 5 contacts.</div>')

CHEVRON = ('<svg class="chev" viewBox="0 0 10 24" aria-hidden="true"><path d="M2 3 8 12 2 21" fill="none" '
           f'stroke="{B.GRAY}" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>')


def ladder(f, profile_word):
    n = f["contacts"]
    steps = [(f["contacts"], f"All {profile_word}", "100%"),
             (f["corporate"], "Chose a corporate profile", fp(share(f["corporate"], n))),
             (f["corp10k"], "…at a company with 10,000+ employees", fp(share(f["corp10k"], n))),
             (f["strict"], "…with an ICP role: our core ICP", fp(share(f["strict"], n)))]
    cells = []
    for i, (v, lab, sh) in enumerate(steps):
        cls = "step last" if i == len(steps) - 1 else "step"
        cells.append(f'<div class="{cls}"><div class="v">{fn(v)}</div><div class="l">{E(lab)}</div><div class="s">{sh} of all</div></div>')
        if i < len(steps) - 1:
            cells.append(CHEVRON)
    split = (f'Of the {fn(f["corp10k"])} corporates with 10,000+ employees: <b>{fn(f["strict"])}</b> have an ICP role '
             f'(and no revenue below $100M), <b>{fn(f["c10_other_role"])}</b> have another role or revenue below $100M, and '
             f'<b>{fn(f["c10_no_title"])}</b> left the job title empty, so we can\'t tell.')
    also = (f'Also ICP: <b>{fn(f["gov"])}</b> government contacts ({fp(share(f["gov"], n), 1)}). Secondary ICP: '
            f'<b>{fn(f["founders"])}</b> startup founders or C-level ({fp(share(f["founders"], n), 1)}).')
    return f'<div class="ladder">{"".join(cells)}</div><div class="note">{split}<br>{also}</div>'


def icp_outcomes(d, f):
    rows = {r["ICP result"]: r for r in d["by_icp"]}
    n = f["contacts"]
    out = ['<table class="t"><tr><th>ICP check</th><th>Contacts</th><th>Qualified by sales</th>'
           '<th>Disqualified by sales</th><th>With a deal</th><th>Became a customer</th></tr>']
    for k in ICP_ORDER:
        r = rows[k]
        cls = "icp-row" if k.startswith("ICP") else ""
        out.append(f'<tr class="{cls}"><td>{E(ICP_LABEL[k])}</td><td>{fn(r["Contacts"])}</td>'
                   f'<td>{fp(r.get("Qualified %"), 1)}</td><td>{fp(r.get("Disqualified %"), 1)}</td>'
                   f'<td>{fp(r.get("With deal %"), 1)}</td><td>{fp(r.get("Customer %"), 1)}</td></tr>')
    out.append(f'<tr class="tot"><td>All contacts</td><td>{fn(n)}</td><td>{fp(share(f["qualified"], n), 1)}</td>'
               f'<td>{fp(share(f["disqualified"], n), 1)}</td><td>{fp(share(f["deals"], n), 1)}</td>'
               f'<td>{fp(share(f["customers"], n), 1)}</td></tr></table>')
    return "".join(out)


def bar_list(items, total=None, kind="src", top=None, width=150, label_w=None):
    """items = [(label, value)] -> rows with a bar, the count and its share of total."""
    items = items[:top] if top else items
    mx = max(num(v) or 0 for _, v in items) or 1
    rows = []
    for lab, v in items:
        sh = f'<td class="sh">{fp(share(v, total))}</td>' if total else ""
        rows.append(f'<tr><td class="bl"{(" style=width:" + str(label_w) + "px") if label_w else ""}>{E(lab)}</td>'
                    f'<td class="l">{bar(v, mx, kind, width)}</td><td>{fn(v)}</td>{sh}</tr>')
    return '<table class="t bars">' + "".join(rows) + "</table>"


def persona_section(d, f):
    pers = [(ROLE_LABEL.get(r["Persona"], r["Persona"]), r["Contacts (all profiles)"]) for r in d["persona_all"]]
    titles = [(r["Job title (lower case)"], r["Contacts"]) for r in d["job_titles"]][:12]
    pxs = d["persona_x_size"]
    sizes = ["Large (10,000+)", "Mid (1,000–9,999)", "Small (<1,000)", "Unknown"]
    hdr = "".join(f"<th>{E(s)}</th>" for s in sizes)
    body = []
    for r in pxs:
        p = r["Persona (corporate profile)"]
        cells = []
        for s in sizes:
            hl = ' class="hl"' if s == "Large (10,000+)" and p in ("C-level", "Innovation / R&D / Digital") else ""
            cells.append(f"<td{hl}>{fn(r.get(s))}</td>")
        body.append(f"<tr><td>{E(ROLE_LABEL.get(p, p))}</td>{''.join(cells)}</tr>")
    grid = f'<table class="t"><tr><th>Corporate profile: role</th>{hdr}</tr>{"".join(body)}</table>'
    return pers, titles, grid


def assets_table(c):
    rows = c["assets"]
    out = ['<table class="t tight"><tr><th>Content asset (100+ contacts)</th><th>Contacts</th><th class="l">Corporates</th><th></th>'
           '<th>Corporates 10,000+</th></tr>']
    ranked = sorted(rows, key=lambda r: -(r["Corporate share"] or 0))
    for r in ranked:
        name = re.sub(r"\s*\|\s*Plug and Play\s*$", "", r["Asset"]).replace(" ", " ")
        hl = ' class="hl"' if (r["Corporate 10,000+ share"] or 0) >= 0.10 else ""
        out.append(f'<tr><td>{E(name)}</td><td>{fn(r["Contacts"])}</td><td class="l">{bar(r["Corporate share"], 0.4, "icp", 70)}</td>'
                   f'<td>{fp(r["Corporate share"])}</td><td{hl}>{fp(r["Corporate 10,000+ share"], 1)}</td></tr>')
    return "".join(out) + "</table>"


CSS = B.BASE_CSS + f"""
@page {{ size: A4; margin: 0; }}
* {{ box-sizing: border-box; }}
body {{ font-size: 8.4pt; line-height: 1.36; }}
.page {{ width: 210mm; min-height: 296mm; padding: 9mm 10mm 6mm; display: flex; flex-direction: column; gap: 2.8mm; break-after: page; }}
.page:last-child {{ break-after: auto; }}
.kicker {{ color: {B.GRAY}; font-size: 7.4pt; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; }}
h1 {{ color: {B.DARK_BLUE}; font-size: 19pt; font-weight: 400; margin: 1px 0 0; line-height: 1.12; }}
.lede {{ color: {B.GRAY}; font-size: 8.2pt; margin-top: 2px; }}
.card {{ padding: 3mm 3.8mm; }}
h2 {{ color: {B.DARK_BLUE}; font-size: 10.4pt; margin: 0 0 1.8mm; font-weight: 700; }}
h2 .hint {{ color: {B.GRAY}; font-size: 7.3pt; font-weight: 400; margin-left: 6px; }}
.answer .bar-y {{ width: 34px; height: 4px; border-radius: 2px; background: {B.YELLOW}; margin-bottom: 6px; }}
.answer .h {{ color: {B.DARK_BLUE}; font-size: 12.2pt; font-weight: 700; line-height: 1.28; }}
.answer .s {{ margin-top: 3px; font-size: 8.7pt; }}
.tiles {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 2.6mm; }}
.tile {{ padding: 2.6mm 3.2mm; }}
.tile .l {{ color: {B.GRAY}; font-size: 6.9pt; font-weight: 700; letter-spacing: .06em; text-transform: uppercase; }}
.tile .v {{ color: {B.DARK_BLUE}; font-size: 17pt; font-weight: 700; line-height: 1.15; margin-top: 2px; }}
.tile .v small {{ font-size: 9pt; color: {B.GRAY}; }}
.tile .c {{ color: {B.TEXT}; font-size: 7.6pt; margin-top: 2px; }}
.cols {{ display: grid; grid-template-columns: 1fr 1fr; gap: 2.8mm; }}
.cols3 {{ display: grid; grid-template-columns: 1.15fr 1fr; gap: 2.8mm; }}
.top-line {{ width: 26px; height: 3px; border-radius: 2px; margin-bottom: 4px; }}
ul.list {{ list-style: none; padding: 0; margin: 0; }}
ul.list li {{ position: relative; padding-left: 13px; margin: 0 0 1.7mm; font-size: 8.2pt; }}
ul.list li:last-child {{ margin-bottom: 0; }}
ul.list li::before {{ content: ""; position: absolute; left: 1px; top: 5px; width: 6px; height: 6px; border-radius: 50%; background: {B.DARK_BLUE}; }}
ul.list.good li::before {{ background: {B.GOOD}; }} ul.list.bad li::before {{ background: {B.BAD}; }}
ul.list b, .note b {{ color: {B.DARK_BLUE}; }}
ol.steps {{ margin: 0; padding-left: 16px; }}
ol.steps li {{ margin: 0 0 1.6mm; font-size: 8.2pt; padding-left: 2px; }}
ol.steps li::marker {{ color: {B.DARK_BLUE}; font-weight: 700; }}
.icpbox {{ font-size: 7.9pt; }}
.icpbox b {{ color: {B.DARK_BLUE}; }}
table.t {{ width: 100%; border-collapse: collapse; }}
table.t th {{ text-align: right; color: {B.GRAY}; font-weight: 700; font-size: 6.8pt; border-bottom: 1.3px solid {B.DARK_BLUE}; padding: 2px 4px 3px; white-space: nowrap; }}
table.t th:first-child, table.t td:first-child {{ text-align: left; }}
table.t th.l, table.t td.l {{ text-align: left; width: 1%; }}
table.t td {{ border-bottom: 1px solid {B.TINT}; padding: 2px 4px; font-size: 7.9pt; text-align: right; white-space: nowrap; }}
table.t tr.tot td {{ font-weight: 700; color: {B.DARK_BLUE}; border-top: 1.3px solid {B.DARK_BLUE}; border-bottom: none; }}
table.t tr.icp-row td {{ background: {B.BG}; font-weight: 700; color: {B.DARK_BLUE}; }}
table.t td.good {{ color: {B.GOOD}; font-weight: 700; }} table.t td.bad {{ color: {B.BAD}; font-weight: 700; }}
table.t td.hl {{ background: {B.TINT}; font-weight: 700; color: {B.DARK_BLUE}; }}
table.bars td {{ border-bottom: none; padding: 1.2px 4px; }}
table.bars td.bl {{ white-space: nowrap; }}
table.bars td.sh {{ color: {B.GRAY}; width: 1%; }}
.bar {{ display: inline-block; height: 7px; border-radius: 4px; background: {B.TINT}; vertical-align: middle; position: relative; }}
.bar i {{ position: absolute; left: 0; top: 0; bottom: 0; border-radius: 4px; background: {B.DARK_BLUE}; }}
.bar.sec i {{ background: {B.BLUE}; }} .bar.oth i {{ background: {B.BLUE_LIGHT}; }} .bar.src i {{ background: {B.DARK_BLUE}; }}
.tag {{ display: inline-block; font-size: 6.3pt; font-weight: 700; padding: 0 5px; border-radius: 7px; margin-left: 3px; vertical-align: 1px; }}
.tag.icp {{ background: {B.DARK_BLUE}; color: #fff; }} .tag.sec {{ background: {B.TINT}; color: {B.DARK_BLUE}; }}
.ladder {{ display: flex; align-items: stretch; gap: 1.5mm; }}
.ladder .step {{ flex: 1; background: {B.BG}; border-radius: {B.RADIUS}; padding: 2.4mm 3mm; }}
.ladder .step.last {{ background: {B.DARK_BLUE}; }}
.ladder .step .v {{ color: {B.DARK_BLUE}; font-size: 15pt; font-weight: 700; line-height: 1.1; }}
.ladder .step .l {{ font-size: 7.6pt; margin-top: 1px; }}
.ladder .step .s {{ color: {B.GRAY}; font-size: 7pt; }}
.ladder .step.last .v, .ladder .step.last .l {{ color: #fff; }} .ladder .step.last .s {{ color: {B.TINT_2}; }}
.chev {{ width: 9px; height: 24px; align-self: center; flex: none; }}
.note {{ color: {B.TEXT}; font-size: 7.8pt; margin-top: 2mm; }}
.legend {{ color: {B.GRAY}; font-size: 6.9pt; margin-top: 1.6mm; }}
.so {{ font-size: 8.2pt; margin: 0 0 2mm; }}
.so b {{ color: {B.DARK_BLUE}; }}
.foot {{ margin-top: auto; color: {B.GRAY}; font-size: 6.8pt; display: flex; justify-content: space-between; }}
.method {{ color: {B.GRAY}; font-size: 6.8pt; line-height: 1.34; columns: 2; column-gap: 6mm; }}
table.t.tight td {{ padding: 1.3px 4px; }}
.method p {{ margin: 0 0 1.4mm; break-inside: avoid; }}
.method b {{ color: {B.TEXT}; }}
"""


def page(inner, n, total, title=""):
    return (f'<section class="page">{inner}<div class="foot"><span>Plug and Play · ICP vs conversions · '
            f'HubSpot, {E(D["meta"]["period"])} · aggregated data only</span><span>{n} / {total}</span></div></section>')


def head(kicker, title, lede=""):
    return (f'<div><div class="kicker">{E(kicker)}</div><h1>{E(title)}</h1>'
            + (f'<div class="lede">{lede}</div>' if lede else "") + "</div>")


def build_html():
    TOTAL = 6
    jn, cn = JF["contacts"], CF["contacts"]
    db = F["db"]
    mk = db["marketing_contacts"]
    # ---------------- page 1: the answer
    ratio_j = JF["c10_deal"] / JF["deal_rate"]
    ratio_c = CF["c10_deal"] / CF["deal_rate"]
    p1 = head("Plug and Play · HubSpot analysis · Raúl's exercise", "Who converts, and are they our ICP?",
              "The marketing email database and the two ways people convert, checked against the 2024 ICP. "
              "<b>JTP Contact us</b>: type of profile, original source and job title. <b>Content downloads</b>: RFI type of "
              "profile and original source.")
    p1 += (f'<div class="card answer"><div class="bar-y"></div><div class="h">Most people who convert are not our ICP: '
           f'{fp(share(JF["startup"], jn))} of JTP contacts are startups, and corporates with 10,000+ employees are only '
           f'{fp(share(JF["corp10k"], jn), 1)} of JTP and {fp(share(CF["corp10k"], cn), 1)} of content contacts.</div>'
           f'<div class="s">The ones we do reach convert: corporates with 10,000+ employees get a deal about '
           f'{ratio_j:.0f} times more often than the average JTP contact ({fp(JF["c10_deal"], 1)} vs {fp(JF["deal_rate"], 1)}) '
           f'and {ratio_c:.0f} times more often among content downloaders ({fp(CF["c10_deal"], 1)} vs {fp(CF["deal_rate"], 1)}).</div></div>')
    p1 += ('<div class="tiles">'
           f'<div class="card tile"><div class="l">People who converted</div><div class="v">{fn(F["unique"])}</div>'
           f'<div class="c">{fn(jn)} JTP Contact us · {fn(cn)} content downloads · {fn(F["both"])} did both</div></div>'
           f'<div class="card tile"><div class="l">Corporates 10,000+</div><div class="v">{fp(share(JF["corp10k"], jn), 1)} <small>· {fp(share(CF["corp10k"], cn), 1)}</small></div>'
           f'<div class="c">of JTP · of content contacts ({fn(JF["corp10k"])} · {fn(CF["corp10k"])} people)</div></div>'
           f'<div class="card tile"><div class="l">Startups in JTP Contact us</div><div class="v">{fp(share(JF["startup"], jn))}</div>'
           f'<div class="c">{fn(JF["startup"])} of {fn(jn)} chose "startup"</div></div>'
           f'<div class="card tile"><div class="l">No job title</div><div class="v">{fp(share(JF["no_title"], jn))} <small>· {fp(share(CF["no_title"], cn))}</small></div>'
           f'<div class="c">of JTP · of content contacts: the ICP role can\'t be checked</div></div></div>')
    found = [
        f'<b>The email database is large and mostly unknown.</b> {fn(mk)} marketing contacts; {fp(share(db["marketing_no_engagement_12m"], mk))} '
        f'opened or clicked no email in 12 months. Only {fp(share(db["answered"]["Type of profile form"], db["all_contacts"]), 1)} of all contacts '
        'answered the profile question.',
        f'<b>JTP Contact us is a startup door.</b> {fn(JF["startup"])} of {fn(jn)} contacts ({fp(share(JF["startup"], jn))}) chose startup; '
        f'corporates are {fp(share(JF["corporate"], jn))} and governments {fp(share(JF["gov"], jn), 1)}.',
        f'<b>Content reaches more corporates, plus a lot of noise.</b> {fp(share(CF["corporate"], cn))} chose corporate, but '
        f'{fn(CF["noise"])} ({fp(share(CF["noise"], cn))}) chose Other, Job seeker or Student.',
        f'<b>Paid social brings volume, few buyers.</b> {fp(share(CF["paid_social"]["Contacts"], cn))} of content contacts came from paid social, '
        f'with {words(CF["paid_social"]["With a deal"])} deals; {fp(JF["src"]["Paid social"]["Disqualified %"], 1)} of JTP paid-social contacts were disqualified by sales.',
    ]
    working = [
        f'<b>Large corporates convert.</b> JTP: {fp(JF["c10_qual"], 1)} qualified by sales and {fp(JF["c10_deal"], 1)} with a deal. '
        f'Content: {fn(CF["c10_customers"])} became customers.',
        f'<b>Core ICP from content:</b> {fn(CF["icp"]["ICP – Corporate"]["With a deal"])} of {fn(CF["strict"])} have a deal and '
        f'{fn(CF["icp"]["ICP – Corporate"]["Became customer"])} became customers.',
        f'<b>Governments engage.</b> JTP: {fp(JF["prof"]["Government"]["Qualified %"], 1)} qualified by sales, '
        f'{fp(JF["prof"]["Government"]["With deal %"], 1)} with a deal.',
        f'<b>Offline sources bring buyers</b> (imports, manual entries, integrations). Content contacts from them: '
        f'{fp(CF["offline"]["With deal %"], 1)} with a deal, {fn(CF["offline"]["Became customer"])} customers.',
    ]
    p1 += (f'<div class="cols"><div class="card"><div class="top-line" style="background:{B.BAD}"></div><h2>What\'s flagged</h2>'
           f'<ul class="list bad">{"".join("<li>" + x + "</li>" for x in found)}</ul></div>'
           f'<div class="card"><div class="top-line" style="background:{B.GOOD}"></div><h2>What\'s working</h2>'
           f'<ul class="list good">{"".join("<li>" + x + "</li>" for x in working)}</ul></div></div>')
    steps = [
        "<b>Ask for job title and company size</b> on every JTP and content form (required, or a role dropdown), so each lead can be checked against the ICP.",
        f'<b>Keep non-ICP contacts away from sales</b>: route Other, Job seeker and Student ({fp(share(CF["noise"], cn))} of content contacts) to nurture.',
        f'<b>Point paid social at the ICP</b>: it brings {fp(share(CF["paid_social"]["Contacts"], cn))} of content contacts but '
        f'{words(CF["paid_social"]["With a deal"])} deals, so target corporate innovation and C-level roles and judge it on qualified leads, not downloads.',
        "<b>Score leads with the ICP</b>: align HubSpot's MQL fit score with the 2024 ICP and fill the empty Persona and Decision Maker Fit fields.",
        "<b>Promote the content that attracts corporates</b> in corporate campaigns (page 6).",
    ]
    p1 += (f'<div class="cols3"><div class="card"><div class="top-line" style="background:{B.DARK_BLUE}"></div><h2>What to do next</h2>'
           f'<ol class="steps">{"".join("<li>" + x + "</li>" for x in steps)}</ol></div>'
           f'<div class="card icpbox"><div class="top-line" style="background:{B.YELLOW}"></div><h2>The ICP used (Buyer Persona 2024)</h2>'
           '<p style="margin:0 0 1.4mm"><b>Core ICP:</b> corporates with 10,000+ employees and $100M+ revenue (where known), in an '
           'innovation, R&amp;D, digital or C-level role. <b>Also ICP:</b> governments and public institutions. <b>Secondary:</b> startup founders.</p>'
           '<p style="margin:0">Outcome = Lead Status (SFDC) and deals in HubSpot, not lifecycle stage: '
           '95% of JTP contacts are set to "Sales Qualified Lead", so it doesn\'t show quality.</p></div></div>')
    pages = [page(p1, 1, TOTAL)]

    # ---------------- page 2: the marketing email database
    srcs = sorted(db["options"]["Original Source"], key=lambda o: -o["contacts"])
    tpf = {}
    for grp in ("Type of profile form (visible)", "Type of profile form (hidden, legacy)"):
        for o in db["options"][grp]:
            lab = o["option"]
            key = ("Startup" if "startup" in lab.lower() else "Corporate" if ("corporation" in lab.lower()) else
                   "Investor or VC" if "investor" in lab.lower() else "Government" if "government" in lab.lower() else
                   "University" if "university" in lab.lower() else "Student" if "student" in lab.lower() else
                   "Other" if lab == "Other" else "Legacy or blank values")
            tpf[key] = tpf.get(key, 0) + o["contacts"]
    tpf_items = sorted(tpf.items(), key=lambda kv: -kv[1])
    rfi_map = {"RFI I work for a major corporation": "Corporate", "Other": "Other", "RFI I work at a world class startup company": "Startup",
               "RFI I am an angel investor or venture capitalist": "Investor or VC", "I am looking for job opportunities": "Job seeker",
               "I am a student": "Student", "RFI I work at a university": "University", "RFI I work for a government": "Government"}
    rfi_items = [(rfi_map.get(o["option"], o["option"]), o["contacts"]) for o in db["options"]["RFI Type of profile"]]
    ans = db["answered"]
    allc = db["all_contacts"]
    p2 = head("1 · The marketing email database", "Who is in the database we email",
              'All HubSpot contacts on 8 Oct 2026. Marketing contacts are the ones we can email and are billed for.')
    p2 += ('<div class="tiles">'
           f'<div class="card tile"><div class="l">All contacts</div><div class="v">{fn(allc)}</div><div class="c">in HubSpot</div></div>'
           f'<div class="card tile"><div class="l">Marketing contacts</div><div class="v">{fn(mk)}</div>'
           f'<div class="c">{fp(share(mk, db["marketing_tier"]))} of the {fn(db["marketing_tier"])} paid tier</div></div>'
           f'<div class="card tile"><div class="l">Not engaged in 12 months</div><div class="v">{fp(share(db["marketing_no_engagement_12m"], mk))}</div>'
           f'<div class="c">{fn(db["marketing_no_engagement_12m"])} marketing contacts with no open or click</div></div>'
           f'<div class="card tile"><div class="l">From offline sources</div><div class="v">{fp(share(F["db_offline"], F["db_sources_total"]))}</div>'
           f'<div class="c">events, imports and lists ({fn(F["db_offline"])} contacts)</div></div></div>')
    known = [("Job title", ans["Job title"]), ("Type of profile form (JTP forms)", ans["Type of profile form"]),
             ("RFI Type of profile (content forms)", ans["RFI Type of profile"]), ("Ebook Type of Profile (older forms)", ans["Ebook Type of Profile"])]
    p2 += (f'<div class="cols"><div class="card"><h2>We know little about most contacts</h2>'
           f'<p class="so">Job title is filled for {fp(share(ans["Job title"], allc))} of all contacts and the profile question for '
           f'{fp(share(ans["Type of profile form"], allc), 1)}, so most of the database can\'t be checked against the ICP.</p>'
           f'{bar_list(known, allc, "src", width=70)}</div>'
           f'<div class="card"><h2>Original source<span class="hint">how each contact first reached HubSpot</span></h2>'
           f'{bar_list([(o["option"], o["contacts"]) for o in srcs], F["db_sources_total"], "src", width=95)}</div></div>')
    p2 += (f'<div class="cols"><div class="card"><h2>Type of profile form<span class="hint">{fn(ans["Type of profile form"])} contacts answered</span></h2>'
           f'{bar_list(tpf_items, ans["Type of profile form"], "src", width=95)}</div>'
           f'<div class="card"><h2>RFI Type of profile<span class="hint">{fn(ans["RFI Type of profile"])} contacts answered</span></h2>'
           f'{bar_list(rfi_items, ans["RFI Type of profile"], "src", width=95)}</div></div>')
    def pair(label, jv, cv, rate_only=False):
        if rate_only:
            return f'<tr><td>{E(label)}</td><td></td><td>{jv}</td><td></td><td>{cv}</td></tr>'
        return (f'<tr><td>{E(label)}</td><td>{fn(jv)}</td><td>{fp(share(jv, jn), 1)}</td>'
                f'<td>{fn(cv)}</td><td>{fp(share(cv, cn), 1)}</td></tr>')
    side = [pair("Chose a corporate profile", JF["corporate"], CF["corporate"]),
            pair("Corporates with 10,000+ employees", JF["corp10k"], CF["corp10k"]),
            pair("Core ICP (10,000+ employees, ICP role, $100M+)", JF["strict"], CF["strict"]),
            pair("Government (ICP)", JF["gov"], CF["gov"]),
            pair("Startups (secondary ICP)", JF["startup"], CF["startup"]),
            pair("No job title", JF["no_title"], CF["no_title"]),
            pair("Qualified by sales", JF["qualified"], CF["qualified"]),
            pair("With a deal", JF["deals"], CF["deals"]),
            pair("Became a customer", JF["customers"], CF["customers"])]
    p2 += (f'<div class="card"><h2>The two ways people convert, side by side<span class="hint">{E(D["meta"]["period"])}</span></h2>'
           '<table class="t"><tr><th></th><th>JTP Contact us</th><th>Share</th><th>Content downloads</th><th>Share</th></tr>'
           f'<tr><td>Forms · submissions</td><td>{fn(JF["forms"])} · {fn(JF["submissions"])}</td><td></td>'
           f'<td>{fn(CF["forms"])} · {fn(CF["submissions"])}</td><td></td></tr>'
           f'<tr class="icp-row"><td>Contacts (people)</td><td>{fn(jn)}</td><td>100%</td><td>{fn(cn)}</td><td>100%</td></tr>'
           + pair("…of them marketing contacts (we can email them)", JF["marketing"], CF["marketing"])
           + "".join(side) + "</table>"
           f'<div class="legend">{fn(F["both"])} people did both: {fn(F["unique"])} people in total, '
           f'{fp(share(F["unique"], allc), 1)} of all contacts. Shares are of each group\'s contacts.</div></div>')
    pages.append(page(p2, 2, TOTAL))

    # ---------------- page 3: JTP – type of profile
    p3 = head("2 · JTP Contact us · type of profile", "Who uses the JTP contact forms",
              f'{fn(jn)} contacts from {fn(JF["submissions"])} submissions on {fn(JF["forms"])} JTP forms. Profile = the "Type of profile form" answer.')
    p3 += (f'<div class="card"><h2>By type of profile<span class="hint">rates = share of each profile\'s contacts</span></h2>'
           f'<p class="so"><b>Startups are {fp(share(JF["startup"], jn))} of contacts</b> but rarely become deals '
           f'({fp(JF["prof"]["Startup"]["With deal %"], 2)}). Corporates with 10,000+ employees and governments are few, '
           f'and they convert best.</p>{profile_table(J, JF)}'
           f'{LEGEND}</div>')
    p3 += (f'<div class="card"><h2>How many fit the ICP<span class="hint">each step keeps the contacts that pass</span></h2>'
           f'{ladder(JF, "JTP contacts")}</div>')
    p3 += f'<div class="card"><h2>What happens to each ICP group</h2>{icp_outcomes(J, JF)}</div>'
    pages.append(page(p3, 3, TOTAL))

    # ---------------- page 4: JTP – source and job title
    pers, titles, grid = persona_section(J, JF)
    src = JF["src"]
    p4 = head("3 · JTP Contact us · original source and job title", "Where they come from, and what they do")
    p4 += (f'<div class="card"><h2>By original source</h2><p class="so"><b>Organic search and direct traffic bring '
           f'{fp(share(src["Organic search"]["Contacts"] + src["Direct traffic"]["Contacts"], jn))} of JTP contacts.</b> '
           f'Paid social ({fp(src["Paid social"]["Disqualified %"], 1)}) and paid search ({fp(src["Paid search"]["Disqualified %"], 1)}) '
           f'have the most contacts disqualified by sales; offline sources have the best deal rate ({fp(src["Offline sources"]["With deal %"], 1)}).</p>'
           f'{source_table(J, JF)}</div>')
    p4 += (f'<div class="cols"><div class="card"><h2>By job title<span class="hint">grouped by role</span></h2>'
           f'<p class="so"><b>{fp(share(JF["no_title"], jn))} gave no job title.</b> Most of the titles given are C-level or founder; '
           f'innovation, R&amp;D or digital roles are {fn(next(v for p, v in pers if p.startswith("Innovation")))} contacts.</p>'
           f'{bar_list(pers, jn, "src", width=70)}</div>'
           f'<div class="card"><h2>Most common job titles</h2>{bar_list(titles, None, "src", width=80)}</div></div>')
    large = {r["Persona (corporate profile)"]: r["Large (10,000+)"] for r in J["persona_x_size"]}
    icp_role = [num(large[p]) for p in ("C-level", "Innovation / R&D / Digital")]
    shaded = (f'Shaded: an ICP role at a 10,000+ company ({fn(sum(icp_role))} contacts: the {fn(JF["strict"])} core ICP on page 3, '
              'plus a few with revenue below $100M). ' if None not in icp_role else 'Shaded: an ICP role at a 10,000+ company. ')
    p4 += (f'<div class="card"><h2>Corporate profile: role × company size<span class="hint">{fn(JF["corporate"])} corporate contacts</span></h2>'
           f'<p class="so">{shaded}At 10,000+ companies the biggest group left the job title empty '
           f'({fn(large["Empty"])} of {fn(JF["corp10k"])}).</p>{grid}</div>')
    pages.append(page(p4, 4, TOTAL))

    # ---------------- page 5: content – type of profile and source
    csrc = CF["src"]
    p5 = head("4 · Content downloads · RFI type of profile and original source", "Who downloads our content",
              f'{fn(cn)} contacts from {fn(CF["submissions"])} downloads on {fn(CF["forms"])} forms. Profile = the "RFI Type of profile" answer.')
    p5 += (f'<div class="card"><h2>By RFI type of profile</h2><p class="so"><b>{fp(share(CF["corporate"], cn))} chose corporate</b>, '
           f'twice JTP\'s share, but {fp(share(CF["noise"], cn))} chose Other, Job seeker or Student. Corporates with 10,000+ employees '
           f'get a deal {fp(CF["c10_deal"], 1)} of the time.</p>{profile_table(C, CF)}{LEGEND}</div>')
    p5 += f'<div class="card"><h2>How many fit the ICP</h2>{ladder(CF, "content contacts")}</div>'
    p5 += (f'<div class="card"><h2>By original source</h2><p class="so"><b>Paid social brings {fp(share(csrc["Paid social"]["Contacts"], cn))} '
           f'of downloaders</b> and {words(csrc["Paid social"]["With a deal"])} deals. Offline sources (imports, manual entries, integrations) bring '
           f'{fn(csrc["Offline sources"]["Contacts"])} contacts and {fn(csrc["Offline sources"]["Became customer"])} customers.</p>'
           f'{source_table(C, CF)}</div>')
    pages.append(page(p5, 5, TOTAL))

    # ---------------- page 6: content assets, ICP outcomes, method
    m = D["method"]
    p6 = head("5 · Content downloads · what attracts corporates", "Which content brings corporates")
    p6 += (f'<div class="card"><h2>Assets with 100+ contacts, by share of corporates<span class="hint">shaded = 10%+ corporates with 10,000+ employees</span></h2>'
           f'{assets_table(C)}</div>')
    p6 += f'<div class="card"><h2>What happens to each ICP group</h2>{icp_outcomes(C, CF)}</div>'
    method = [
        ("Data", "HubSpot read-only pull on 8 Oct 2026; form submissions 1 Oct 2025 – 30 Sep 2026. JTP = 8 JTP contact forms "
                 "(Startups, Corporates, Governments, Universities, VCs, APAC, All profiles, Contact us). Content = 66 CDW, RFI and resource forms."),
        ("Unit", "Unique contacts: someone who submitted several times counts once, with the profile and outcome stored on the contact."),
        ("Profile", "JTP: \"Type of profile form\". Content: \"RFI Type of profile\" (or \"Type of profile form\", then \"Ebook Type of Profile\" "
                    "when empty, mainly Japan forms)."),
        ("Role", "From the job title: innovation, R&D or digital; C-level; founder; government; other; empty."),
        ("Company size", "Employees on the contact, then its company; if empty, the size dropdowns."),
        ("Outcome", "Lead Status (SFDC): qualified or converted; disqualified. Deal = at least one associated deal. "
                    "Customer = entered the Customer stage."),
        ("Database", "All contacts, marketing contacts and engagement: HubSpot searches and the billing page, 8 Oct 2026. "
                     "\"Not engaged\" = no open and no click in 12 months, never-engaged included."),
        ("Not included", "1,733 submissions on the two \"All profiles – dependent fields\" forms have no contact record in HubSpot."),
        ("GDPR", "Aggregated counts only; counts of 1–4 shown as \"<5\" and rates based on them as \"–\". No names, emails, "
                 "companies or record IDs. Raw data was processed in a temporary workspace and deleted. Tables: ICP_vs_Conversions_2026.xlsx."),
    ]
    p6 += ('<div class="card"><h2>How this was measured</h2><div class="method">'
           + "".join(f"<p><b>{E(k)}.</b> {E(v)}</p>" for k, v in method) + "</div></div>")
    pages.append(page(p6, 6, TOTAL))
    assert m  # method text from the data file is mirrored in the workbook
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>ICP vs Conversions</title>"
            f"<style>{CSS}</style></head><body>{''.join(pages)}</body></html>")


# ------------------------------------------------------------------ workbook (Google Sheets ready: values only, header rows frozen)
def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    FONT = "Trebuchet MS"
    blue = B.DARK_BLUE.lstrip("#")
    title_f = Font(name=FONT, size=14, bold=True, color=blue)
    note_f = Font(name=FONT, size=9, italic=True, color=B.GRAY.lstrip("#"))
    head_f = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    body_f = Font(name=FONT, size=10, color=B.TEXT.lstrip("#"))
    bold_f = Font(name=FONT, size=10, bold=True, color=blue)
    head_fill = PatternFill("solid", fgColor=blue)
    plain = {"Qualified / converted (sales)": "Qualified by sales", "Disqualified (sales)": "Disqualified by sales",
             "Became customer": "Became a customer", "Strict ICP (corporate or government)": "ICP: core corporate or government",
             "Marketing contacts": "Marketing contacts (can be emailed)", "With deal %": "With a deal %",
             "Customer %": "Became a customer %", "Persona (corporate profile)": "Corporate profile: role",
             "Persona assigned": "Role group", "Job title (lower case)": "Job title",
             "Persona": "Role group (from job title)", "ICP result": "ICP check",
             "ICP result vs HubSpot MQL fit score": "ICP check vs HubSpot MQL fit score",
             "ICP result vs HubSpot MQL fit score CDW": "ICP check vs HubSpot MQL fit score (CDW)"}
    values = {"Profile": PROFILE_LABEL, "ICP result": ICP_LABEL, "ICP result vs HubSpot MQL fit score": ICP_LABEL,
              "ICP result vs HubSpot MQL fit score CDW": ICP_LABEL, "Persona": ROLE_LABEL, "Persona (corporate profile)": ROLE_LABEL,
              "Persona assigned": ROLE_LABEL,
              "": {"Strict ICP – Corporate (role + 10,000+)": "Core ICP: corporate, ICP role, 10,000+ employees, $100M+ (where known)",
                   "Startup profile (secondary ICP group)": "Startup profile (secondary ICP)"}}

    def sheet(name, title, note, blocks, widths=None):
        ws = wb.create_sheet(name)
        ws["A1"], ws["A2"] = title, note
        ws["A1"].font, ws["A2"].font = title_f, note_f
        r = 4
        first_header = None
        for rows, label_key in blocks:
            if not rows:
                continue
            cols = list(rows[0].keys())
            for ci, c in enumerate(cols, 1):
                cell = ws.cell(row=r, column=ci, value=plain.get(c, c))
                cell.font, cell.fill = head_f, head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            first_header = first_header or r
            for row in rows:
                r += 1
                for ci, c in enumerate(cols, 1):
                    v = row[c]
                    if isinstance(v, str):
                        v = values.get(c, {}).get(v, v)
                    cell = ws.cell(row=r, column=ci, value=v)
                    cell.font = body_f
                    if isinstance(v, float) and ("%" in c or "share" in c.lower()):
                        cell.number_format = "0.0%"
                    elif isinstance(v, (int, float)):
                        cell.number_format = "#,##0"
                    if isinstance(v, str) and v in ("<5", "–"):
                        cell.alignment = Alignment(horizontal="right")
            r += 2
        for ci in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(ci)].width = (widths or {}).get(ci, 16 if ci > 1 else 44)
        if first_header:
            ws.freeze_panes = ws.cell(row=first_header + 1, column=1)
        return ws

    def summary_rows():
        keys = list(S.keys())
        return [{"": k, "JTP Contact us": S[k]["jtp"], "Content downloads": S[k]["content"]} for k in keys]

    ws = sheet("Summary", "ICP vs conversions – summary",
               f"HubSpot, {D['meta']['period']}. Aggregated counts only; counts of 1–4 shown as '<5'. No personal data.",
               [(summary_rows(), "")], {1: 46, 2: 20, 3: 20})
    db = F["db"]
    dbrows = [{"Measure": "All contacts", "Contacts": db["all_contacts"]},
              {"Measure": "Marketing contacts (can be emailed, billed)", "Contacts": db["marketing_contacts"]},
              {"Measure": "Marketing-contact tier", "Contacts": db["marketing_tier"]},
              {"Measure": "Marketing contacts with no open or click in 12 months", "Contacts": db["marketing_no_engagement_12m"]},
              {"Measure": "Marketing contacts that never opened a marketing email", "Contacts": db["marketing_never_opened"]}]
    dbrows += [{"Measure": f"Contacts with a value: {k}", "Contacts": v} for k, v in db["answered"].items()]
    srcrows = [{"Original source (all contacts)": o["option"], "Contacts": o["contacts"]} for o in db["options"]["Original Source"]]
    prof_rows = [{"Profile question": grp, "Answer": o["option"], "Contacts": o["contacts"]}
                 for grp in ("Type of profile form (visible)", "Type of profile form (hidden, legacy)", "RFI Type of profile", "Ebook Type of Profile")
                 for o in db["options"][grp]]
    sheet("Database", "The marketing email database", "All HubSpot contacts on 8 Oct 2026 (searches and the billing page). "
          "Profile answers and original source are counted over all contacts.", [(dbrows, ""), (srcrows, ""), (prof_rows, "")],
          {1: 52, 2: 44, 3: 14})
    for grp, d, label in (("JTP", J, "JTP Contact us"), ("Content", C, "Content downloads")):
        sheet(f"{grp} – profile", f"{label} – by type of profile",
              "Unique contacts. Outcome = Lead Status (SFDC); deals = associated deals. Rates as a share of each row's contacts.",
              [(d["by_profile"], "Profile")], {1: 36})
        sheet(f"{grp} – source", f"{label} – by original source", "Original source of the contact (HubSpot).",
              [(d["by_source"], "Original source")], {1: 24})
        sheet(f"{grp} – job title", f"{label} – job title", "Role groups from the job title; corporate profile by role and company size; "
              "titles held by 5+ contacts.", [(d["persona_all"], ""), (d["persona_x_size"], ""), (d["size_all"], ""), (d["job_titles"], "")],
              {1: 40})
        sheet(f"{grp} – ICP", f"{label} – ICP check and outcome", "ICP check per contact (see Method). Last table: HubSpot's MQL fit "
              "score for each ICP group.", [(d["by_icp"], "ICP result"), (d["fit_score"], "")], {1: 40})
        sheet(f"{grp} – forms", f"{label} – forms", "Form submissions and unique contacts per form.", [(d["forms"], "Form")], {1: 56})
    sheet("Content – assets", "Content downloads – assets with 100+ contacts", "Unique contacts per asset (page title).",
          [(C["assets"], "Asset")], {1: 64})
    mrows = [{"Topic": k, "Definition": v} for k, v in D["method"].items()]
    ws = sheet("Method", "Method and definitions", "How the numbers were built.", [(mrows, "")], {1: 22, 2: 120})
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for name in wb.sheetnames:
        for row in wb[name].iter_rows(min_row=5, max_col=1):
            for cell in row:
                if cell.value in ("All contacts", "Contacts", "Total"):
                    cell.font = bold_f
    wb.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    html_path = os.path.join(OUT, STEM + ".html")
    open(html_path, "w", encoding="utf-8").write(build_html())
    build_xlsx(os.path.join(OUT, STEM + ".xlsx"))
    print("written", html_path, "and", STEM + ".xlsx")
    if "--pdf" in sys.argv:
        font = B.font_in_use()
        pdf = os.path.join(OUT, STEM + ".pdf")
        subprocess.run([CHROMIUM, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf}", "file://" + os.path.abspath(html_path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
        pages = re.search(r"Pages:\s+(\d+)", info).group(1)
        fonts = subprocess.run(["pdffonts", pdf], capture_output=True, text=True).stdout
        names = {re.sub(r"^[A-Z]{6}\+", "", ln.split()[0]) for ln in fonts.splitlines()[2:] if ln.strip()}
        other = sorted(n for n in names if font.replace(" ", "").lower() not in n.lower())
        print(f"PDF: {pdf} · {pages} pages · font {font}" + (f" · WARNING non-brand fonts: {other}" if other else ""))


if __name__ == "__main__":
    main()
